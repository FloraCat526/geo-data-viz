#!/usr/bin/env python3
"""Profile local geographic data without network calls or coordinate conversion."""
import argparse
import csv
import datetime as dt
import hashlib
import io
import json
import math
import re
import sys
from collections import Counter
from pathlib import Path

MISSING = {"", "null", "none", "nan", "na", "n/a", "-"}
ALIASES = {
    "lng": {"lng", "lon", "long", "longitude", "经度", "经度坐标"},
    "lat": {"lat", "latitude", "纬度", "纬度坐标"},
    "label": {"name", "label", "title", "名称", "名字", "地点名称", "城市", "city"},
    "category": {"category", "type", "group", "类别", "分类", "类型", "分组"},
    "time": {"time", "timestamp", "date", "datetime", "时间", "日期", "采集时间"},
    "address": {"address", "addr", "地址", "详细地址", "location", "地点", "城市", "city"},
}
GEO_TYPES = {"Point", "MultiPoint", "LineString", "MultiLineString", "Polygon", "MultiPolygon", "GeometryCollection"}


def key_name(value):
    return re.sub(r"[\s_\-]+", "", str(value)).lower()


def missing(value):
    return value is None or (isinstance(value, str) and value.strip().lower() in MISSING)


def number(value):
    if missing(value) or isinstance(value, bool) or isinstance(value, (list, dict)):
        return None
    try:
        result = float(value.strip() if isinstance(value, str) else value)
        return result if math.isfinite(result) else None
    except (ValueError, TypeError, OverflowError):
        return None


def safe_json(value, issues, path=""):
    """Represent invalid JSON nonfinite values as null and retain an audit warning."""
    if isinstance(value, float) and not math.isfinite(value):
        issues.append(path or "root")
        return None
    if isinstance(value, dict):
        return {str(k): safe_json(v, issues, f"{path}.{k}") for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [safe_json(v, issues, f"{path}[{i}]") for i, v in enumerate(value)]
    if isinstance(value, (dt.datetime, dt.date)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def quantiles(values):
    vals = sorted(values)
    if not vals:
        return None
    def q(p):
        position = (len(vals) - 1) * p
        lo, hi = math.floor(position), math.ceil(position)
        return vals[lo] + (vals[hi] - vals[lo]) * (position - lo)
    return {"min": vals[0], "p25": q(.25), "p50": q(.5), "p75": q(.75), "p95": q(.95), "max": vals[-1]}


def parse_time(value):
    if isinstance(value, (dt.datetime, dt.date)):
        value = value.isoformat()
    if not isinstance(value, str) or missing(value):
        return None
    value = value.strip()
    try:
        if re.match(r"^\d{4}-\d{2}-\d{2}(?:$|[T ])", value):
            parsed = dt.datetime.fromisoformat(value.replace("Z", "+00:00"))
            return parsed.astimezone(dt.timezone.utc).isoformat().replace("+00:00", "Z") if parsed.tzinfo else parsed.isoformat()
        if re.fullmatch(r"\d{4}/\d{1,2}/\d{1,2}", value):
            return dt.datetime.strptime(value, "%Y/%m/%d").date().isoformat()
    except ValueError:
        pass
    return None


def load_input(path, sheet=None):
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        try:
            import openpyxl
        except ImportError as exc:
            raise ValueError("Excel requires optional openpyxl; export CSV or use an environment with openpyxl installed.") from exc
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if sheet and sheet not in book.sheetnames:
            raise ValueError(f"Sheet {sheet!r} does not exist. Available: {book.sheetnames}")
        tab = book[sheet] if sheet else book.active
        all_rows = tab.iter_rows(values_only=True)
        header = next(all_rows, None)
        if not header:
            raise ValueError("Excel sheet is empty.")
        headers = [str(h).strip() if h is not None else f"column_{i+1}" for i, h in enumerate(header)]
        check_headers(headers)
        rows = [dict(zip(headers, row)) for row in all_rows if any(v is not None for v in row)]
        selected = tab.title
        sheets = book.sheetnames
        book.close()
        return rows, None, {"format": "excel", "sheet": selected, "availableSheets": sheets, "rowNumberMeaning": "data record index starting at 1; blank Excel rows skipped"}
    if suffix == ".xls":
        raise ValueError("Legacy .xls is unsupported; export .xlsx or CSV first.")
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("Input is neither UTF-8 nor GB18030. Export a UTF-8 file.")
    if suffix in {".json", ".geojson"} or text.lstrip().startswith(("{", "[")):
        root = json.loads(text)
        if isinstance(root, dict) and root.get("type") == "FeatureCollection":
            if not isinstance(root.get("features"), list):
                raise ValueError("FeatureCollection.features must be an array.")
            return root["features"], root, {"format": "geojson", "encoding": encoding}
        if isinstance(root, dict) and (root.get("type") == "Feature" or root.get("type") in GEO_TYPES):
            feat = root if root["type"] == "Feature" else {"type": "Feature", "geometry": root, "properties": {}}
            return [feat], root, {"format": "geojson", "encoding": encoding}
        records = root
        if isinstance(root, dict):
            keys = [k for k in ("records", "data") if isinstance(root.get(k), list)]
            if len(keys) != 1:
                raise ValueError("JSON must be records[], {records:[]}, {data:[]}, or GeoJSON; ambiguous roots require explicit extraction.")
            records = root[keys[0]]
        if not isinstance(records, list):
            raise ValueError("JSON records must be an array.")
        return records, None, {"format": "json-records", "encoding": encoding}
    delimiter = "\t" if suffix == ".tsv" else ","
    if suffix not in {".csv", ".tsv"}:
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",\t;").delimiter
        except csv.Error:
            pass
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    if not reader.fieldnames:
        raise ValueError("Input has no header.")
    headers = [h.strip() for h in reader.fieldnames]
    check_headers(headers)
    reader.fieldnames = headers
    rows = []
    for row in reader:
        if None in row:
            row["__parse_error"] = "more cells than header fields"
            row.pop(None)
        rows.append(row)
    return rows, None, {"format": "tsv" if delimiter == "\t" else "csv", "encoding": encoding, "rowNumberMeaning": "data record index starting at 1; header excluded"}


def check_headers(headers):
    if len(set(headers)) != len(headers) or any(not h for h in headers):
        raise ValueError("Headers must be nonempty and unique after trimming; rename duplicate/empty columns before profiling.")


def declared_crs(root):
    if not isinstance(root, dict):
        return None
    metadata = root.get("metadata")
    if isinstance(metadata, dict) and metadata.get("dataCrs"):
        return metadata["dataCrs"]
    return root.get("crs")


def normalize_crs(value):
    if value is None:
        return None
    compact = re.sub(r"[^a-z0-9]", "", json.dumps(value).lower())
    if "gcj02" in compact or "gcj2" in compact:
        return "gcj02"
    if "bd09" in compact or "bd9" in compact:
        return "bd09"
    if "wgs84" in compact or "crs84" in compact or "epsg4326" in compact or "epsg64326" in compact:
        return "wgs84"
    return "unknown"


def id_like(field, values):
    name = key_name(field)
    if name in {"id", "fid", "uuid", "index", "序号", "编号", "编码", "代码", "邮编", "zipcode", "postalcode", "adcode"}:
        return True
    if re.search(r"(?:^|[_\-\s])(id|code)$", str(field).lower()) or re.search(r"(?:Id|ID|Code)$", str(field)) or name.endswith(("编号", "编码", "代码", "邮编")):
        return True
    strings = [v.strip() for v in values if isinstance(v, str) and not missing(v)]
    return bool(strings) and sum(bool(re.match(r"^0\d+$", v)) for v in strings) / len(strings) > .5


def inspect_fields(records):
    fields = sorted({str(k) for row in records if isinstance(row, dict) for k in row})
    stats = {}
    candidates = {role: [] for role in (*ALIASES, "value")}
    for field in fields:
        values = [row.get(field) for row in records if isinstance(row, dict)]
        present = [v for v in values if not missing(v)]
        nums = [n for v in values if (n := number(v)) is not None]
        unique = Counter(json.dumps(v, ensure_ascii=False, sort_keys=True, default=str) for v in present)
        identifier = id_like(field, values)
        stat = {"nonMissing": len(present), "missing": len(values)-len(present), "numeric": len(nums), "numericInvalid": len(present)-len(nums), "unique": len(unique), "idLike": identifier, "quantiles": quantiles(nums)}
        if len(unique) <= 50:
            stat["topValues"] = [{"value": json.loads(v), "count": c} for v, c in unique.most_common(10)]
        normalized = key_name(field)
        coordinate = False
        for role, aliases in ALIASES.items():
            if normalized in aliases:
                candidates[role].append(field)
                coordinate = coordinate or role in {"lng", "lat"}
        if present and len(nums)/len(present) >= .8 and not identifier and not coordinate and normalized not in ALIASES["time"]:
            candidates["value"].append(field)
        times = [t for v in present if (t := parse_time(v)) is not None]
        if times:
            utc = sorted(t for t in times if t.endswith("Z"))
            local = sorted(t for t in times if not t.endswith("Z"))
            stat["time"] = {"parseable": len(times), "unparseable": len(present)-len(times), "examples": times[:3], "utcRange": [utc[0], utc[-1]] if utc else None, "noTimezoneRange": [local[0], local[-1]] if local else None}
            if len(times)/len(present) >= .8 and field not in candidates["time"]:
                candidates["time"].append(field)
        if present and 2 <= len(unique) <= min(30, max(2, len(present)//2)) and not nums and field not in candidates["category"]:
            candidates["category"].append(field)
        stats[field] = stat
    return stats, candidates


def validate_geometry(geometry):
    """Validate basic geometry structure; does not certify polygon topology."""
    points, warnings = [], []
    if not isinstance(geometry, dict) or geometry.get("type") not in GEO_TYPES:
        raise ValueError("missing or unsupported geometry type")
    def position(p, path):
        if not isinstance(p, list) or len(p) < 2:
            raise ValueError(f"{path}: position must contain at least longitude, latitude")
        if any(isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v) for v in p):
            raise ValueError(f"{path}: all ordinates must be finite numbers")
        lng, lat = p[:2]
        if not (-180 <= lng <= 180 and -90 <= lat <= 90):
            hint = "; possible reversed longitude/latitude" if -180 <= lat <= 180 and -90 <= lng <= 90 else ""
            raise ValueError(f"{path}: longitude/latitude out of range{hint}")
        points.append([lng, lat])
    def line(coords, path, minimum=2, ring=False):
        if not isinstance(coords, list) or len(coords) < minimum:
            raise ValueError(f"{path}: requires at least {minimum} positions")
        for i, p in enumerate(coords):
            position(p, f"{path}[{i}]")
        if ring and coords[0] != coords[-1]:
            raise ValueError(f"{path}: polygon ring must be closed")
        if ring and len({tuple(p[:2]) for p in coords[:-1]}) < 3:
            raise ValueError(f"{path}: polygon ring needs at least 3 distinct positions")
    def polygon(coords, path):
        if not isinstance(coords, list) or not coords:
            raise ValueError(f"{path}: polygon needs an exterior ring")
        for i, ring in enumerate(coords):
            line(ring, f"{path}[{i}]", 4, True)
    def visit(g, path):
        if not isinstance(g, dict):
            raise ValueError(f"{path}: geometry must be an object")
        kind, coords = g.get("type"), g.get("coordinates")
        if kind == "Point":
            position(coords, path)
        elif kind == "LineString":
            line(coords, path)
        elif kind == "Polygon":
            polygon(coords, path)
        elif kind in {"MultiPoint", "MultiLineString", "MultiPolygon"}:
            if not isinstance(coords, list) or not coords:
                raise ValueError(f"{path}: {kind} must be nonempty")
            fn = {"MultiPoint": position, "MultiLineString": line, "MultiPolygon": polygon}[kind]
            for i, child in enumerate(coords):
                fn(child, f"{path}[{i}]")
        elif kind == "GeometryCollection":
            children = g.get("geometries")
            if not isinstance(children, list) or not children:
                raise ValueError(f"{path}: GeometryCollection must be nonempty")
            for i, child in enumerate(children):
                visit(child, f"{path}.geometries[{i}]")
        else:
            raise ValueError(f"{path}: unsupported geometry {kind!r}")
    visit(geometry, "geometry")
    return points, warnings


def spatial_summary(points):
    if not points:
        return {"bounds": None, "wrappedBounds": None, "center": None, "crossesAntimeridian": False}
    lngs = sorted(set(p[0] for p in points))
    lats = [p[1] for p in points]
    gaps = [(lngs[(i+1) % len(lngs)] + (360 if i == len(lngs)-1 else 0) - lngs[i], i) for i in range(len(lngs))]
    largest, i = max(gaps)
    west, east = lngs[(i+1) % len(lngs)], lngs[i]
    span = 360-largest
    center_lng = (west+span/2+180) % 360-180
    return {"bounds": [min(lngs), min(lats), max(lngs), max(lats)], "wrappedBounds": [west, min(lats), east, max(lats)], "crossesAntimeridian": west > east, "longitudeSpan": span, "center": [center_lng, (min(lats)+max(lats))/2], "zeroZeroVertices": sum(p == [0, 0] for p in points), "mercatorClippedVertices": sum(abs(p[1]) > 85.051129 for p in points)}


def area_km2(geometry):
    def ring_area(ring):
        total = 0
        for p, q in zip(ring, ring[1:]):
            delta = (math.radians(q[0]-p[0])+math.pi) % (2*math.pi)-math.pi
            total += delta*(2+math.sin(math.radians(p[1]))+math.sin(math.radians(q[1])))
        return abs(total)*6371.0088**2/2
    def polygon_area(coords):
        return max(0, ring_area(coords[0])-sum(ring_area(ring) for ring in coords[1:]))
    kind = geometry["type"]
    if kind == "Polygon":
        return polygon_area(geometry["coordinates"])
    if kind == "MultiPolygon":
        return sum(polygon_area(p) for p in geometry["coordinates"])
    if kind == "GeometryCollection":
        return sum(area_km2(g) for g in geometry["geometries"])
    return 0


def geometry_parts(geometry):
    """Yield atomic geometry parts without changing source geometry or feature IDs."""
    kind = geometry["type"]
    if kind == "GeometryCollection":
        for child in geometry["geometries"]:
            yield from geometry_parts(child)
    elif kind.startswith("Multi"):
        for coordinates in geometry["coordinates"]:
            yield {"type": kind[5:], "coordinates": coordinates}
    else:
        yield geometry


def visualization_candidates(features, roles, part_counts, duplicate_locations):
    """Unranked structural candidates; the agent still decides intent and units."""
    result = []
    def add(kind, status, reason, requires=()):
        result.append({"type": kind, "status": status, "available": status == "eligible",
                       "requires": list(requires), "reason": reason})
    def metric_state(kind):
        relevant = [f for f in features if any(g["type"] == kind for g in geometry_parts(f["geometry"]))]
        values = [f["properties"]["__viz"]["value"] for f in relevant]
        nums = [v for v in values if v is not None]
        if not roles["value"]:
            return "needs-transform", ["value-field"], nums
        if not nums:
            return "blocked", ["finite-metric-on-valid-geometry"], nums
        return "eligible", [], nums
    if part_counts["Point"]:
        add("points", "eligible", "Equal-size locations; no numeric metric required. This list is not a ranking.")
        status, requires, values = metric_state("Point")
        add("point-color", status, "Map a confirmed intensity, rate or signed change to color; units determine the scale.", requires)
        if values and min(values) < 0:
            status, requires = "needs-transform", ["explicit-absolute-value-with-sign-encoding"]
        add("bubble", status, "Area requires a nonnegative magnitude. Signed data needs an explicit transform and a separate sign channel; missing values stay separate.", requires)
        add("cluster", "needs-transform", "Repeated point locations detected; retain access to all entities." if duplicate_locations else "Assess overlap at the intended viewport; feature count alone does not determine suitability.", ["confirm-point-grain-and-view-overlap"])
        add("heatmap", "needs-transform", "Choose count density or valid nonnegative contribution weights explicitly; temperature and rates are not additive contributions.", ["confirm-count-grain-or-nonnegative-contribution"])
        add("breathing-points", "needs-transform", "Optional emphasis for selected events or the user's visual preference; not a default analytical recommendation.", ["confirm-emphasis-meaning"])
    if part_counts["LineString"]:
        add("route-lines", "eligible", "Render existing vertices; distinguish actual paths from OD relations and do not invent direction or timestamps.")
    if part_counts["Polygon"]:
        add("area-outline", "eligible", "Display known boundaries without inventing a metric.")
        status, requires, _ = metric_state("Polygon")
        add("choropleth", status, "Use a meaningful area measure; rates require denominator-aware aggregation, and null remains missing.", requires)
    for role in ("category", "time"):
        if roles[role]:
            add(f"{role}-filter", "eligible", "Use the bound source field; a time field alone does not imply a trajectory.")
            result[-1]["field"] = roles[role]
    return result


def profile(path, out_dir, crs=None, fields=None, sheet=None):
    rows, root, source = load_input(path, sheet)
    source["fileName"] = path.name
    source["sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
    warnings = []
    props = [r.get("properties") or {} if source["format"] == "geojson" and isinstance(r, dict) else r for r in rows]
    stats, candidates = inspect_fields(props)
    roles = {}
    fields = fields or {}
    for role in ("lng", "lat", "value", "label", "category", "time"):
        override = fields.get(role)
        if override is not None and override not in stats:
            raise ValueError(f"Unknown --{role}-field {override!r}; available fields: {list(stats)}")
        roles[role] = override if override is not None else (candidates[role][0] if role != "value" and len(candidates[role]) == 1 else None)
    if roles["lng"] and roles["lng"] == roles["lat"]:
        raise ValueError("Longitude and latitude must use different fields.")
    candidates["value"] = [f for f in candidates["value"] if f not in {roles["lng"], roles["lat"], roles["time"]}]
    if candidates["value"] and not roles["value"]:
        warnings.append("numeric_value_unbound: select a meaningful metric with --value-field; identifiers are excluded from suggestions")
    declared = declared_crs(root)
    data_crs = crs or (normalize_crs(declared) if declared is not None else ("wgs84" if root is not None else "unknown"))
    if declared is not None:
        source["declaredCrs"] = declared
        if normalize_crs(declared) != "wgs84":
            warnings.append("non_rfc7946_crs: source declares a non-WGS84 or unrecognized CRS; source coordinates retained")
        if crs and normalize_crs(declared) != crs:
            warnings.append("crs_override: --crs overrides the source declaration; no coordinates were converted")
    if data_crs == "unknown":
        warnings.append("unknown_crs: numeric coordinate ranges cannot identify WGS84, GCJ-02 or BD-09; establish source provenance")
    invalid, features, all_points = [], [], []
    geom_counts, seen_records, seen_points, seen_ids = Counter(), Counter(), Counter(), Counter()
    part_counts = Counter()
    coordinate_missing = root is None and not (roles["lng"] and roles["lat"])
    for index, row in enumerate(rows, 1):
        reasons = []
        if not isinstance(row, dict):
            invalid.append({"row": index, "reasons": ["record must be an object"]})
            continue
        original_props = row.get("properties") if root is not None else row
        if original_props is None:
            original_props = {}
        if not isinstance(original_props, dict):
            invalid.append({"row": index, "reasons": ["properties must be an object or null"]})
            continue
        nonfinite = []
        properties = safe_json(original_props, nonfinite)
        if nonfinite:
            warnings.append(f"nonfinite_properties: row {index} converted nonfinite values to null at {', '.join(nonfinite[:4])}")
        canonical_row = json.dumps(safe_json(row, []), sort_keys=True, ensure_ascii=False, separators=(",", ":"))
        row_hash = hashlib.sha256(canonical_row.encode("utf-8")).hexdigest()[:20]
        seen_records[row_hash] += 1
        if root is not None:
            if row.get("type") != "Feature":
                reasons.append("FeatureCollection item must have type Feature")
            geometry = row.get("geometry")
        else:
            geometry = None
            if row.get("__parse_error"):
                reasons.append(row["__parse_error"])
            if coordinate_missing:
                reasons.append("longitude/latitude fields unavailable or ambiguous")
            else:
                lng, lat = number(row.get(roles["lng"])), number(row.get(roles["lat"]))
                for role, value in (("lng", lng), ("lat", lat)):
                    if value is None:
                        raw = row.get(roles[role])
                        reasons.append(f"{role}: {'missing' if missing(raw) else 'not a finite number'}")
                if lng is not None and lat is not None:
                    geometry = {"type": "Point", "coordinates": [lng, lat]}
        points = []
        if not reasons:
            try:
                points, geometry_warnings = validate_geometry(geometry)
                warnings.extend(geometry_warnings)
            except ValueError as exc:
                reasons.append(str(exc))
        if reasons:
            invalid.append({"row": index, "sourceId": row.get("id"), "reasons": reasons})
            continue
        label = properties.get(roles["label"]) if roles["label"] else None
        category = properties.get(roles["category"]) if roles["category"] else None
        time_value = properties.get(roles["time"]) if roles["time"] else None
        if "__viz" in properties:
            backup = "__source_viz"
            while backup in properties:
                backup = "_" + backup
            properties[backup] = properties["__viz"]
        if root is not None and "id" in row:
            backup = "__source_feature_id"
            while backup in properties:
                backup = "_" + backup
            properties[backup] = safe_json(row["id"], [])
        properties["__viz"] = {"value": number(properties.get(roles["value"])) if roles["value"] else None, "label": str(label) if not missing(label) else f"Record {index}", "category": str(category) if not missing(category) else None, "time": parse_time(time_value)}
        fid = f"f_{row_hash}"
        seen_ids[fid] += 1
        if seen_ids[fid] > 1:
            fid += f"_{seen_ids[fid]}"
        features.append({"type": "Feature", "id": fid, "geometry": geometry, "properties": properties})
        geom_counts[geometry["type"]] += 1
        all_points.extend(points)
        for part in geometry_parts(geometry):
            part_counts[part["type"]] += 1
            if part["type"] == "Point":
                seen_points[tuple(part["coordinates"][:2])] += 1
    if not rows:
        status = "needs_coordinates"
        warnings.append("empty_dataset: no records found")
    elif coordinate_missing:
        status = "needs_geocoding" if candidates["address"] else "needs_coordinates"
    elif not features:
        status = "needs_coordinates"
    elif data_crs == "unknown":
        status = "needs_crs"
    else:
        status = "ready"
    spatial = spatial_summary(all_points)
    if spatial.get("zeroZeroVertices"):
        warnings.append("zero_zero: (0,0) was retained; verify whether it represents a real location or a missing-value sentinel")
    if spatial.get("mercatorClippedVertices"):
        warnings.append("mercator_latitude_limit: valid polar coordinates may be clipped by Web Mercator basemaps")
    if spatial["crossesAntimeridian"]:
        warnings.append("antimeridian: use wrappedBounds or explicit longitude wrapping; naive bounds will zoom to most of the world")
    if roles["value"]:
        value_stats = stats[roles["value"]]
        if value_stats["missing"] or value_stats["numericInvalid"]:
            warnings.append("metric_gaps: missing/nonfinite/nonnumeric values remain null and must not be mapped to zero")
        if value_stats["quantiles"] and value_stats["quantiles"]["min"] < 0:
            warnings.append("signed_metric: use diverging color or explicit positive/negative channels; negative values cannot define radius")
    if roles["time"]:
        times = [f["properties"]["__viz"]["time"] for f in features]
        if any(t is None for t in times):
            warnings.append("time_gaps: unparseable or missing timestamps remain null; epoch units and ambiguous dates need an explicit parser")
    duplicate_locations = sum(n-1 for n in seen_points.values())
    recommendations = visualization_candidates(features, roles, part_counts, duplicate_locations)
    if duplicate_locations:
        warnings.append("overlapping_locations: multiple point parts share coordinates; preserve entity access and confirm counting grain before aggregation")
    if part_counts["Polygon"]:
        warnings.append("polygon_validation_scope: ring structure and coordinate values checked; self-intersection/topology and ring orientation not certified")
        spatial["approxPolygonAreaKm2"] = sum(area_km2(f["geometry"]) for f in features) if data_crs == "wgs84" else None
        spatial["areaMethod"] = "spherical estimate, Earth radius 6371.0088 km; holes subtracted; not suitable for legal measurement" if data_crs == "wgs84" else "not calculated for unverified/non-WGS84 datum"
    warnings = list(dict.fromkeys(warnings))
    counts = {"inputRecords": len(rows), "validFeatures": len(features), "invalidRecords": len(invalid), "geometryTypes": dict(geom_counts), "geometryParts": dict(part_counts), "vertices": len(all_points), "duplicateRecords": sum(n-1 for n in seen_records.values()), "duplicatePointLocations": duplicate_locations, "uniquePointLocations": len(seen_points)}
    metadata = {"schemaVersion": 1, "status": status, "dataCrs": data_crs, "roles": roles, "warnings": warnings, "source": source}
    result = {**metadata, "roleCandidates": candidates, "counts": counts, "invalidRows": invalid, "fields": stats, "spatial": spatial, "recommendations": recommendations}
    collection = {"type": "FeatureCollection", "metadata": metadata, "features": features}
    out_dir.mkdir(parents=True, exist_ok=True)
    for name in ("profile.json", "normalized.geojson"):
        if (out_dir / name).resolve() == path.resolve():
            raise ValueError("Output path would overwrite the input file; choose another --out-dir.")
    for name, data in (("profile.json", result), ("normalized.geojson", collection)):
        destination = out_dir / name
        destination.write_text(json.dumps(safe_json(data, []), ensure_ascii=False, indent=2, allow_nan=False)+"\n", encoding="utf-8")
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("--out-dir", required=True, type=Path)
    parser.add_argument("--crs", choices=("wgs84", "gcj02", "bd09", "unknown"), default=None, help="Source CRS; never performs coordinate conversion. Unspecified tables remain unknown.")
    parser.add_argument("--sheet", help="Excel worksheet name; defaults to active sheet and records available sheets.")
    for role in ("lng", "lat", "value", "label", "category", "time"):
        parser.add_argument(f"--{role}-field")
    args = parser.parse_args()
    try:
        result = profile(args.input, args.out_dir, args.crs, {role: getattr(args, f"{role}_field") for role in ("lng", "lat", "value", "label", "category", "time")}, args.sheet)
    except (ValueError, OSError, json.JSONDecodeError, csv.Error) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    print(json.dumps({"status": result["status"], "dataCrs": result["dataCrs"], "counts": result["counts"], "outputs": [str(args.out_dir / name) for name in ("profile.json", "normalized.geojson")]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
